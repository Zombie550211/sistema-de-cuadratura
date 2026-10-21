"""RRHH corrige y elimina marcas, borra importaciones y empleados; Excel y locales."""

import io

from openpyxl import Workbook

from tests.test_api import ATTLOG, _employee, _upload


def _punches(client, **extra):
    params = {"date_from": "2026-09-25", "date_to": "2026-09-25", **extra}
    return client.get("/api/attendance/punches", params=params).json()["items"]


def _xlsx(rows) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_edit_and_delete_punch_survive_reimport(login):
    client = login("rrhh")
    assert _upload(client).status_code == 201
    first = next(p for p in _punches(client) if p["badge"] == "36" and p["punched_at"].startswith("2026-09-25T07"))

    r = client.patch(f"/api/attendance/punches/{first['id']}", json={"punched_at": "2026-09-25T07:00:00-06:00"})
    assert r.status_code == 200, r.text
    assert r.json()["original_punched_at"] is not None

    other = next(p for p in _punches(client) if p["badge"] == "36" and p["id"] != first["id"])
    assert client.delete(f"/api/attendance/punches/{other['id']}").status_code == 204

    # Una descarga nueva del reloj trae las mismas marcas: no deben volver ni duplicarse.
    r = _upload(client, ATTLOG + b"            36\t2026-09-26 07:00:00\t1\t4\t1\t0\r\n", "BINH221860214_attlog(2).dat")
    assert r.status_code == 201 and r.json()["inserted"] == 1
    badge36 = [p for p in _punches(client) if p["badge"] == "36"]
    assert [p["punched_at"][11:16] for p in badge36] == ["07:00"]


def test_delete_import_removes_its_punches_and_allows_reupload(login):
    client = login("rrhh")
    batch = _upload(client).json()
    assert client.delete(f"/api/attendance/imports/{batch['id']}").status_code == 204
    assert _punches(client) == []
    assert _upload(client).status_code == 201


def test_delete_employee_keeps_punches_without_name(login):
    client = login("rrhh")
    emp = _employee(client, badge_number="83")
    _upload(client)
    assert any(p["employee"] for p in _punches(client))
    assert client.delete(f"/api/employees/{emp['id']}").status_code == 204
    assert client.get(f"/api/employees/{emp['id']}").status_code == 404
    assert all(p["employee"] is None for p in _punches(client, badge="83"))


def test_excel_detects_local_by_fingerprints_and_links_names(login):
    client = login("rrhh")
    ana = _employee(client, fingerprints={"LA18": "5"})
    _employee(client, code="E-002", first_name="Luis", last_name="Pérez", fingerprints={"LA18": "6"})
    content = _xlsx([
        ["Reporte de marcaciones"],
        ["No.", "Nombre", "Fecha", "Hora"],
        [5, "Ana López", "25/09/2026", "07:01:00"],
        [6, "Luis Pérez", "25/09/2026", "7:15 a. m."],
        [5, "Ana López", "25/09/2026", "12:00:00"],
    ])
    r = client.post("/api/attendance/imports", files={"file": ("marcas.xlsx", content, "application/octet-stream")})
    assert r.status_code == 201, r.text
    assert r.json()["inserted"] == 3 and r.json()["unknown_badges"] == 0

    locales = client.get("/api/attendance/locales").json()
    device = next(d for d in locales["devices"] if d["fingerprint_key"] == "LA18")
    assert device["punches"] == 3
    punches = _punches(client, device_id=device["id"])
    assert {p["employee"]["id"] for p in punches if p["badge"] == "5"} == {ana["id"]}
    assert all(p["employee"] for p in punches)


def test_excel_unknown_local_asks_which_one(login):
    client = login("rrhh")
    content = _xlsx([["ID", "Fecha/Hora"], [99, "2026-09-25 08:00:00"]])
    r = client.post("/api/attendance/imports", files={"file": ("x.xlsx", content, "application/octet-stream")})
    assert r.status_code == 422 and "local" in r.json()["detail"]
    r = client.post("/api/attendance/imports", files={"file": ("x.xlsx", content, "application/octet-stream")}, data={"local": "LB8"})
    assert r.status_code == 201, r.text
    assert r.json()["device"]["fingerprint_key"] == "LB8"


def test_link_badge_assigns_fingerprint_for_that_local(login):
    client = login("rrhh")
    _upload(client)
    emp = _employee(client)
    device_id = _punches(client)[0]["device_id"]
    r = client.post("/api/attendance/badges/link", json={"device_id": device_id, "badge": "36", "employee_id": emp["id"]})
    assert r.status_code == 204, r.text
    assert {p["employee"]["id"] for p in _punches(client, badge="36")} == {emp["id"]}


def test_roster_excel_to_text(login):
    client = login("rrhh")
    content = _xlsx([["N°", "EMPLEADO"], [1, "ANA LOPEZ"]])
    r = client.post("/api/employees/import/excel", files={"file": ("p.xlsx", content, "application/octet-stream")})
    assert r.status_code == 200 and r.json()["text"] == "N°\tEMPLEADO\n1\tANA LOPEZ"


def test_punches_by_section(login):
    client = login("rrhh")
    _employee(client, badge_number="83", area="ventas")
    _employee(client, code="E-002", first_name="Luis", last_name="Pérez", badge_number="36", area="operativo")
    _upload(client)
    assert {p["badge"] for p in _punches(client, area="ventas")} == {"83"}
    assert {p["badge"] for p in _punches(client, area="operativo")} == {"36"}
    assert _punches(client, area="administrativo") == []
    locales = client.get("/api/attendance/locales", params={"area": "operativo"}).json()["devices"]
    assert [d["punches"] for d in locales] == [2]


def test_excel_reads_every_punch_sheet(login):
    """Como el Excel de RRHH: ASISTENCIA trae solo la primera marca del día y BASE todas."""
    client = login("rrhh")
    _employee(client, badge_number="59", fingerprints={"LA17": "59"})
    wb = Workbook()
    wb.active.title = "ASISTENCIA"
    for row in (["N", "EMPLEADO", "FECHA", "HORA"], [59, "Ana López", "29/09/2026", "07:20:22"]):
        wb.active.append(row)
    base = wb.create_sheet("BASE")
    for row in ([None, "N ", "FECHA", "HORA"], [None, 59, "29/09/2026", "07:20:22"], [None, 59, "29/09/2026", "13:09:50"]):
        base.append(row)
    other = wb.create_sheet("EMPLEADOS")
    other.append([59, "ANA LOPEZ"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/attendance/imports", files={"file": ("asistencia.xlsx", buf.getvalue(), "application/octet-stream")}, data={"local": "LA17"})
    assert r.status_code == 201, r.text
    assert r.json()["inserted"] == 2


def test_saved_cut_dates_history(login):
    client = login("rrhh")
    a = {"date_from": "2026-09-01", "date_to": "2026-09-20"}
    b = {"date_from": "2026-09-16", "date_to": "2026-10-05"}
    assert client.put("/api/attendance/cut-date", json=a).status_code == 200
    assert client.put("/api/attendance/cut-date", json=b).status_code == 200
    assert client.put("/api/attendance/cut-date", json=a).status_code == 200
    # La más reciente primero y sin repetidas.
    assert client.get("/api/attendance/cut-dates").json() == [a, b]

    # Quitar la aplicada solo deja de aplicarla: sigue guardada para volver a elegirla.
    assert client.delete("/api/attendance/cut-date").status_code == 204
    assert client.get("/api/attendance/cut-date").json() is None
    assert client.get("/api/attendance/cut-dates").json() == [a, b]

    # Borrarla de las guardadas la quita de la lista (y de la aplicada si lo era).
    assert client.put("/api/attendance/cut-date", json=b).status_code == 200
    assert client.delete("/api/attendance/cut-dates", params=b).status_code == 204
    assert client.get("/api/attendance/cut-dates").json() == [a]
    assert client.get("/api/attendance/cut-date").json() is None
