"""Lectura de archivos de Excel (.xlsx y .xls).

- `sheet_rows`: filas de la primera hoja con datos, con fechas y horas ya convertidas.
- `sheet_as_text`: la hoja como texto separado por tabulaciones, igual que al copiar y pegar
  desde Excel, para reutilizar la importación de la nómina.
- `parse_excel_punches`: marcaciones exportadas del reloj o de su programa a Excel. Busca las
  columnas por su encabezado (ID / huella, fecha y hora, nombre) y, si no hay encabezados,
  asume el orden del attlog: ID, fecha-hora.
"""

import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

from app.services.zk_attlog import AttlogRecord, ParseResult, normalize_badge

XLSX_MAGIC = b"PK\x03\x04"
XLS_MAGIC = b"\xd0\xcf\x11\xe0"


class ExcelError(ValueError):
    pass


def is_excel(content: bytes) -> bool:
    return content.startswith(XLSX_MAGIC) or content.startswith(XLS_MAGIC)


def all_sheet_rows(content: bytes) -> list[list[list]]:
    """Todas las hojas con datos (la activa primero), sin filas vacías."""
    if content.startswith(XLSX_MAGIC):
        sheets = _xlsx_sheets(content)
    elif content.startswith(XLS_MAGIC):
        sheets = _xls_sheets(content)
    else:
        raise ExcelError("El archivo no es un Excel (.xlsx o .xls)")
    cleaned = [
        [r for r in ([_strip(c) for c in row] for row in sheet) if any(c not in (None, "") for c in r)]
        for sheet in sheets
    ]
    cleaned = [rows for rows in cleaned if rows]
    if not cleaned:
        raise ExcelError("El Excel no tiene datos")
    return cleaned


def sheet_rows(content: bytes) -> list[list]:
    return all_sheet_rows(content)[0]


def _strip(value):
    return value.strip() if isinstance(value, str) else value


def _xlsx_sheets(content: bytes) -> list[list[list]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # archivo dañado o protegido
        raise ExcelError("No se pudo abrir el Excel") from exc
    sheets = [book.active, *[ws for ws in book.worksheets if ws is not book.active]]
    return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in sheets if ws is not None]


def _xls_sheets(content: bytes) -> list[list[list]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise ExcelError("No se pudo abrir el Excel") from exc

    def cell(c):
        if c.ctype == xlrd.XL_CELL_DATE:
            value = xlrd.xldate_as_datetime(c.value, book.datemode)
            # Una celda solo con hora llega como 1899-12-31 HH:MM.
            return value.time() if c.value < 1 else value
        if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
            return None
        return c.value

    return [[[cell(c) for c in sheet.row(i)] for i in range(sheet.nrows)] for sheet in book.sheets()]


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y") if value.time() == time.min else value.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\t", " ").replace("\n", " ")


def sheet_as_text(content: bytes) -> str:
    return "\n".join("\t".join(_text(c) for c in row) for row in sheet_rows(content))


# ---------- Marcaciones en Excel ----------

def _norm(text) -> str:
    plain = unicodedata.normalize("NFD", str(text or "")).encode("ascii", "ignore").decode().upper()
    return " ".join(re.findall(r"[A-Z0-9]+", plain))


ID_HEADERS = {
    "ID", "NO", "AC NO", "ACNO", "HUELLA", "N HUELLA", "NO HUELLA", "NUMERO DE HUELLA", "ID RELOJ",
    "ID DE RELOJ", "ID DE USUARIO", "ID USUARIO", "USER ID", "USERID", "ENROLL NUMBER", "ENROLLNUMBER",
    "PERSONNEL ID", "ID PERSONAL", "ID DE EMPLEADO", "ID EMPLEADO", "EMPLOYEE ID", "BADGE", "PIN",
    "NO DE EMPLEADO", "N", "NUMERO",
}
DATETIME_HEADERS = {
    "FECHA HORA", "FECHAHORA", "FECHA Y HORA", "TIEMPO", "DATETIME", "DATE TIME", "MARCACION",
    "MARCA", "CHECKTIME", "CHECK TIME", "HORA DE MARCACION", "FECHA DE MARCACION", "PUNCH TIME",
    "REGISTRO", "TIME",
}
DATE_HEADERS = {"FECHA", "DATE", "DIA"}
TIME_HEADERS = {"HORA", "TIME", "HORA DE ENTRADA", "ENTRADA"}
NAME_HEADERS = {"NOMBRE", "NAME", "EMPLEADO", "NOMBRE COMPLETO", "NOMBRES", "AGENTE", "FULL NAME"}

DATETIME_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M",
    "%d/%m/%Y %I:%M:%S %p", "%d/%m/%Y %I:%M %p", "%Y-%m-%d %I:%M:%S %p", "%Y-%m-%d %I:%M %p",
)
DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y")
TIME_FORMATS = ("%H:%M:%S", "%H:%M", "%I:%M:%S %p", "%I:%M %p")


def _ampm(text: str) -> str:
    # "6:21 a. m." / "6:21 p.m." → "6:21 AM"
    text = re.sub(r"\s*a\.?\s*m\.?\s*$", " AM", text, flags=re.IGNORECASE)
    return re.sub(r"\s*p\.?\s*m\.?\s*$", " PM", text, flags=re.IGNORECASE)


def _as_datetime(value) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        text = _ampm(" ".join(value.split()))
        for fmt in DATETIME_FORMATS:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
    return None


def _as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _as_time(value) -> time | None:
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, (int, float)) and 0 <= value < 1:  # fracción del día de Excel
        return (datetime.min + timedelta(seconds=round(value * 86400))).time()
    if isinstance(value, str):
        text = _ampm(value.strip())
        for fmt in TIME_FORMATS:
            try:
                return datetime.strptime(text, fmt).time()
            except ValueError:
                continue
    return None


def _badge(value) -> str | None:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value or "").strip()
    return normalize_badge(text) if re.fullmatch(r"[0-9A-Za-z]+", text) else None


def _find_columns(row: list) -> dict | None:
    heads = [_norm(c) for c in row]

    def first(options: set[str], skip: set[int] = frozenset()) -> int | None:
        return next((i for i, h in enumerate(heads) if h in options and i not in skip), None)

    id_col = first(ID_HEADERS)
    if id_col is None:
        return None
    date_col = first(DATE_HEADERS, {id_col})
    time_col = first(TIME_HEADERS, {id_col, date_col} - {None})
    dt_col = first(DATETIME_HEADERS - {"TIME"}, {id_col})
    if dt_col is None and date_col is None:
        dt_col = first({"TIME"}, {id_col})
    if dt_col is None and (date_col is None or time_col is None):
        # Una sola columna "Fecha" puede traer fecha y hora juntas.
        if date_col is None:
            return None
        dt_col, date_col = date_col, None
    return {"id": id_col, "datetime": dt_col, "date": date_col, "time": time_col, "name": first(NAME_HEADERS, {id_col})}


def _punch_sheets(content: bytes) -> list[tuple[list[list], dict, int]]:
    """Las hojas que tienen marcas (encabezado con ID y fecha/hora). Un Excel puede traer una
    hoja ya resumida (p. ej. solo la primera marca del día) y otra con todas las marcas: se
    leen todas y se juntan sin repetir."""
    sheets = all_sheet_rows(content)
    found = []
    for rows in sheets:
        for i, row in enumerate(rows[:15]):
            if cols := _find_columns(row):
                found.append((rows, cols, i + 1))
                break
    if not found:  # sin encabezados: ID, fecha-hora como en el attlog
        found.append((sheets[0], {"id": 0, "datetime": 1, "date": None, "time": None, "name": None}, 0))
    return found


def parse_excel_punches(content: bytes, tz: ZoneInfo) -> tuple[ParseResult, dict[str, str]]:
    """Devuelve las marcas y, si el Excel trae nombres, {id de reloj: nombre}."""
    result = ParseResult()
    names: dict[str, str] = {}
    seen: set[tuple[str, datetime]] = set()
    for rows, cols, start in _punch_sheets(content):
        _parse_sheet(rows, cols, start, tz, result, names, seen)
    result.records.sort(key=lambda r: (r.badge, r.punched_at))
    return result, names


def _parse_sheet(rows, cols, start, tz, result: ParseResult, names: dict[str, str], seen: set) -> None:
    cell = lambda row, i: row[i] if i is not None and i < len(row) else None  # noqa: E731

    for line_no, row in enumerate(rows[start:], start=start + 1):
        result.total_lines += 1
        badge = _badge(cell(row, cols["id"]))
        if cols["datetime"] is not None:
            at = _as_datetime(cell(row, cols["datetime"]))
        else:
            d, t = _as_date(cell(row, cols["date"])), _as_time(cell(row, cols["time"]))
            at = datetime.combine(d, t) if d and t else None
        if not badge or at is None:
            result.errors.append((line_no, "Fila sin ID de reloj o sin fecha y hora válidas"))
            continue
        at = at.replace(tzinfo=tz)
        name = cell(row, cols["name"])
        if name and isinstance(name, str) and badge not in names:
            names[badge] = name
        if (badge, at) in seen:
            result.duplicate_lines += 1
            continue
        seen.add((badge, at))
        result.records.append(AttlogRecord(line_no=line_no, badge=badge, punched_at=at, verify_mode=0, status_code=0, work_code=0))
